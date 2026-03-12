#International Item Management List
from deep_translator import GoogleTranslator
from deep_translator.exceptions import LanguageNotSupportedException
from currency_converter import CurrencyConverter
from docx import Document
from docx.shared import Pt

titulo = "CS50P Final Project"
def main():
    prints_the_final_list(listing())

def validate_currency(curncy):
    c = CurrencyConverter()
    try:
        c.convert(1, curncy, "USD")
        return "ok"
    except ValueError:
        print("Invalid currency")
        return "not ok"
def validate_price(numbr):
    try:
        float(numbr)
        return "ok"
    except ValueError:
        print("Invalid price")
        return "not ok"

###faz a listagem
def listing():
    list_of_list = []

    while True:
        name: str = input("What is the name of the item? ")
        description: str = input("What is the description? ")
        while True:
            currency: str = input("What is the currency?(EX:USD/BRL/EUR) ")
            result = validate_currency(currency)
            if result == "ok":
                break
        while True:
            price: float = input("What is the price? ")
            resultado = validate_price(price)
            if resultado == "ok":
                break
        respective_item = {name: [description, currency, float(price)]}
        list_of_list.append(respective_item)
        new_list: str = (input("Do you want to list another item? (yes/no) "))
        if new_list == "no":
            break

    return list_of_list

###printa a lista
def prints_the_final_list(listafds: list):
    print(titulo)
    e=0
    for item in listafds:
        e+=1
        for key, value in item.items():
            print(f"""
{e}. {key}
{value[0]}
Price: ${float(value[2]):.2f} {value[1]}""")

    make_docx(listafds)
    modifies_the_list(listafds)

###modificação da lista
def modifies_the_list(listafds: list):
    modify_list = input("\nDo you want to modify the list? (yes/no) ")
    if modify_list == "no":
        print(titulo)
        e=0
        for item in listafds:
            e+=1
            for key, value in item.items():
                print(f"""
{e}. {key}
{value[0]}
Price: ${float(value[2]):.2f} {value[1]}""")

        make_docx(listafds)
    else:
        modify_or_translate = input("""
+====================================+
|                                    |
> [1] Modify the items               |
|                                    |
> [2] Translate and change currency  |
|                                    |
> [3] Done                           |
|                                    |
+====================================+
Answer: """)

        if modify_or_translate == "1":
            lista_escolhida_dict = choose_item(listafds)
            for key, value in lista_escolhida_dict.items():
                print("Write in the front of what you want to modify")
                nome_modificado = input(f"Name: {key} -> ")
                descrição_modificado = input(f"Description: {value[0]} -> ")
                moeda_modificado_mentira = " "
                while True:
                    try:
                        valor_modificado = input(f"Price: {float(value[2]):.2f} -> ")
                        float(valor_modificado)
                        break
                    except ValueError:
                        if valor_modificado.strip():
                            print("Invalid price")
                            pass
                        else:
                            break

                lista_dos_bgl_modificado = [nome_modificado,descrição_modificado,moeda_modificado_mentira,valor_modificado]

            ## muda a key(nome) ou os values da lista
            i = -2
            for var in lista_dos_bgl_modificado:
                i+=1
                if var.strip():
                    if i == -1:
                        muda_key = {}
                        for key, value in lista_escolhida_dict.items():
                            muda_key[key] = var
                        for old, new in muda_key.items():
                            lista_escolhida_dict[new] = lista_escolhida_dict.pop(old)
                    else:
                        for key, value in lista_escolhida_dict.items():
                                value[i] = var
            modifies_the_list(listafds)

        elif modify_or_translate == "2":
            translate_change_currency(listafds)
            modifies_the_list(listafds)


        elif modify_or_translate == "3":
            print(titulo)
            e=0
            for item in listafds:
                e+=1
                for key, value in item.items():
                    print(f"""
{e}. {key}
{value[0]}
Price: ${float(value[2]):.2f} {value[1]}""")
            make_docx(listafds)

        else:
            print("Invalid answer!")
            modifies_the_list(listafds)

###pergunta qual vai ser o item da lista
def choose_item(listafds: list):
    print("\nChoose which item:")
    print("+")
    e = 0
    for item in listafds:
        e+=1
        for key, value in item.items():
            print(f"""|
> [{e}] {key}""")
    print("""|
+""")
    resposta_choose = input("Answer: ")
    return listafds[int(resposta_choose)-1]


def validate_language(resposta_lingua):
    try:
        h = "isso é um teste"
        GoogleTranslator(source='auto', target=resposta_lingua).translate(h)
        return "ok"
    except LanguageNotSupportedException:
        print("Invalid language")
        return "not ok"

###traduz e muda
def translate_change_currency(listafds):
    resposta_tcc = input("""
+====================================+
|                                    |
> [1] Translate                      |
|                                    |
> [2] Change currency                |
|                                    |
> [3] Both                           |
|                                    |
+====================================+
Answer: """)

    if resposta_tcc == "1":
        while True:
                resposta_qual_lingua = input("What language do you want it to be?(EX:en/pt/es) ")
                resultado = validate_language(resposta_qual_lingua)
                if resultado == "ok":
                    break

        for lista in listafds:
            for key, value in lista.items():
                value[0] = GoogleTranslator(source='auto', target=resposta_qual_lingua).translate(value[0])
            muda_key = {}
            for key, value in lista.items():
                muda_key[key] = GoogleTranslator(source='auto', target=resposta_qual_lingua).translate(key)
            for old, new in muda_key.items():
                lista[new] = lista.pop(old)

    elif resposta_tcc == "2":
        while True:
            resposta_qual_currency = input("What currency do you want it to be?(EX:USD/BRL/EUR) ")
            result = validate_currency(resposta_qual_currency)
            if result == "ok":
                break
        for lista in listafds:
            for key, value in lista.items():
                value[2] = c.convert(value[2], value[1], resposta_qual_currency)
                value[1] = resposta_qual_currency

    elif resposta_tcc == "3":
        while True:
            resposta_qual_lingua = input("What language do you want it to be?(EX:en/pt/es) ")
            resultado = validate_language(resposta_qual_lingua)
            if resultado == "ok":
                break

        for lista in listafds:
            for key, value in lista.items():
                value[0] = GoogleTranslator(source='auto', target=resposta_qual_lingua).translate(value[0])
            muda_key = {}
            for key, value in lista.items():
                muda_key[key] = GoogleTranslator(source='auto', target=resposta_qual_lingua).translate(key)
            for old, new in muda_key.items():
                lista[new] = lista.pop(old)

        c = CurrencyConverter()
        while True:
            resposta_qual_currency = input("What currency do you want it to be?(EX:USD/BRL/EUR) ")
            result = validate_currency(resposta_qual_currency)
            if result == "ok":
                break
        for lista in listafds:
            for key, value in lista.items():
                value[2] = c.convert(value[2], value[1], resposta_qual_currency)
                value[1] = resposta_qual_currency

    else:
        print("Invalid answer!")
        translate_change_currency(listafds)

    return listafds

###faz a lista vira docx
def make_docx(listafds):
    document = Document()

    document.add_heading(f"{titulo}",0)
    style = document.styles["Normal"]
    font = style.font
    font.size = Pt(15)

    for lista in listafds:
            for key, value in lista.items():
                document.add_paragraph()
                #nome
                ndoc = document.add_paragraph(style="List Number")
                ndoc.add_run(f"{key}").bold = True
                #descp
                ddoc = document.add_paragraph(f"{value[0]}")
                #price e currency
                pcdoc = document.add_paragraph()
                pcdoc.add_run("Price: ").bold = True
                pcdoc.add_run(f"${float(value[2]):.2f} ")
                pcdoc.add_run(f"{value[1]}").italic = True

    document.save("CS50P_Project docx.docx")


if __name__ == "__main__":
    main()
